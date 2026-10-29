"""저장된 공개 예시 산출물의 값으로 앱 화면을 오프라인 캡처한다.

새 네트워크 조회나 회계 수치 재계산 없이 워크북에 보관된 값을 표시한다.
실행: python scripts/capture_app_screens.py <example.xlsx> [output-directory]
"""
import os
import sys
import tempfile
import shutil
import datetime as dt
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from PIL import ImageGrab
from openpyxl import load_workbook
import app
import settings


def backup(path):
    if path.exists():
        folder = Path(settings.CACHE_DIR) / 'ui-backups'
        folder.mkdir(parents=True, exist_ok=True)
        stamp = dt.datetime.fromtimestamp(path.stat().st_mtime).strftime('%Y%m%d_%H%M%S_%f')
        target = folder / f'backup_{stamp}_gallery_{path.name}'
        if not target.exists():
            shutil.copy2(path, target)


def capture(workbook, output):
    source = load_workbook(workbook, read_only=True, data_only=True)
    summary = {r[0]: r[1] for r in source['요약'].iter_rows(min_row=2, values_only=True) if r[0]}
    peer_rows = list(source['피어'].iter_rows(min_row=2, values_only=True))
    peer_codes = {r[1] for r in peer_rows}
    beta = {r[1]: r for r in source['베타'].iter_rows(min_row=2, values_only=True)}
    candidates = []
    saved = {}
    for r in source['후보군'].iter_rows(min_row=2, values_only=True):
        if not r[2]:
            continue
        candidates.append(dict(name=r[0], market=r[1], code=r[2], industry=r[3], products=r[4] or '',
                               listed=r[14] or '', settle_month=r[15] or '', cap_eok=None,
                               hits=[x.strip() for x in (r[17] or '').split(',') if x.strip() and '★' not in x],
                               flags=[r[16]] if r[16] else [], recommended='★추천' in (r[17] or '')))
        saved[r[2]] = dict(selected=r[2] in peer_codes, excluded=bool(r[11]), reason=r[12] or '')
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp, \
         mock.patch.object(app.PeerApp, '_startup', return_value=None), \
         mock.patch.object(settings, 'UI_PREFERENCES_PATH', os.path.join(tmp, 'ui.json')), \
         mock.patch.object(settings, 'load_config', return_value={k: '' for k in settings.KEYS}):
        a = app.PeerApp()
        a.geometry('1480x880+25+25')
        cp, sp = a.candidate_panel, a.summary_panel
        a.sess['as_of'] = summary['기준일']
        a.sess['target']['name'] = summary['대상기업']
        a.sess['candidates'] = saved
        a.candidates = candidates
        a.kind_rows = candidates
        cp.name_var.set(summary['대상기업'])
        cp.as_of_var.set(summary['기준일'])
        target = next(r for r in candidates if r['name'] == summary['대상기업'])
        cp.industry_var.set(target['industry'])
        cp.kw_var.set(', '.join(source['후보군'].cell(1, i).value for i in range(6, 12)))
        cp.target_info.configure(text=f"{target['name']} · {target['code']}\n{target['products']}")
        cp.render()
        sp.src_seg.set('산출(KRX)' if summary['베타 소스'] == '산출' else '한공회 파일')
        sp.empty_label.place_forget()
        for i, r in enumerate(peer_rows):
            b = beta[r[1]]
            def fmt(value, format_string):
                return '' if value is None else format_string.format(value)
            values = dict(name=r[0], code=r[1], close=fmt(r[6], '{:,.0f}'),
                          e=fmt(r[10] / 1e8 if r[10] is not None else None, '{:,.0f}'),
                          d=fmt(r[17] / 1e8 if r[17] is not None else None, '{:,.0f}'),
                          de=fmt(r[18], '{:.4f}'), tax=fmt(r[19], '{:.1%}'), raw=fmt(b[2], '{:.4f}'),
                          blume=fmt(r[20], '{:.4f}'), n=fmt(b[4], '{}'), r2=fmt(b[3], '{:.4f}'),
                          kraw=fmt(b[9], '{:.6f}'), kadj=fmt(b[10], '{:.6f}'), kpts=fmt(b[11], '{}'),
                          diff=fmt(b[12], '{:.4f}'), bl=fmt(r[22], '{:.4f}'), bu=fmt(r[23], '{:.4f}'),
                          inc='☑' if r[24] else '☐', debt='상세…', flags=r[25] or '')
            sp.tree.insert('', 'end', iid=r[1], values=[values[k] for k, *_ in sp._COLS],
                           tags=(['odd'] if i % 2 else []) + (['flag'] if r[25] else []))
        sp.agg_label.configure(text=f"βU 평균 {summary['βU 평균(집계)']:.4f} · 목표 D/E {summary['목표 D/E']:.4f} · 재레버 βL {summary['재레버 βL']:.4f} · 집계 {summary['집계 피어 수']}개")
        sp.progress.configure(text=f"저장된 예시 산출물 · {summary['대상기업']} · {summary['기준일']}")
        sp.kicpa_label.configure(text='')
        sp.report_btn.configure(state='normal')
        a.env_label.configure(text='저장 산출물 미리보기')
        a.status_label.configure(text='화면 예시 · 새 조회 없이 기존 엑셀 산출물의 값을 표시했습니다.')
        captures = [('01_candidates.png',False,False), ('02_summary.png',False,True),
                    ('01_candidates_dark.png',True,False), ('02_summary_dark.png',True,True),
                    ('02_loading_contrast.png',False,True)]
        def next_capture(index=0):
            if index == len(captures):
                a.destroy()
                return
            name,dark,is_summary = captures[index]
            a.dark_mode_var.set(dark)
            a._toggle_dark_mode()
            a.tabs.set('2. 요약 · WACC' if is_summary else '1. 후보군')
            if name == '02_loading_contrast.png':
                sp.load_btn.configure(state='disabled')
                sp.report_btn.configure(state='disabled')
                sp.progress.configure(text='조회 단계 표시 예시 · 삼성전자 재무상태표 · 경과시간 표시')
            def save():
                a.lift()
                a.update_idletasks()
                x,y = a.winfo_rootx(),a.winfo_rooty()
                p=output / name
                backup(p)
                ImageGrab.grab(bbox=(x,y,x+a.winfo_width(),y+a.winfo_height())).save(p)
                print(name,flush=True)
                a.after(200,lambda:next_capture(index+1))
            a.after(650,save)
        a.after(500,next_capture)
        a.mainloop()
    source.close()


if __name__ == '__main__':
    capture(sys.argv[1], sys.argv[2] if len(sys.argv)>2 else Path(__file__).resolve().parents[1]/'docs'/'img')
